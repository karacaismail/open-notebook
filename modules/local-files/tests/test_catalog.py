import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'service'))
from catalog import Catalog,query_terms
import pytest

@pytest.fixture
def cat(tmp_path):
    root=tmp_path/'home';root.mkdir();(root/'Downloads').mkdir();(root/'Documents').mkdir()
    return Catalog(tmp_path/'index',{'root':str(root),'exclude':[str(root/'Downloads')],'embedding_model':'embeddinggemma:300m-qat-q8_0'})
def drain(cat):
    for _ in range(100):
        if not cat.process_one():return
    raise AssertionError('Unfinished extraction')
def test_turkish_query_finds_english_names_and_body(cat):
    (cat.root/'Documents/strategy.md').write_text('Strategy roadmap for market expansion.')
    (cat.root/'Documents/other.txt').write_text('Our strategic plan includes customer discovery.')
    cat.scan();drain(cat)
    assert len(cat.search('strateji dökümanım nerede')['results'])==2
    assert query_terms('strategy pdf docx md')[1]==['pdf','docx','md']
    assert len(cat.search('strateji md')['results'])==1

def test_exclusion_symlink_and_prefix_boundary(cat):
    (cat.root/'Downloads/strategy.md').write_text('excluded')
    (cat.root/'Documents/alias.md').symlink_to(cat.root/'Downloads/strategy.md')
    (cat.root/'Documents/aliasdir').symlink_to(cat.root/'Downloads',target_is_directory=True)
    (cat.root/'Downloads-extra').mkdir();(cat.root/'Downloads-extra/strategy.md').write_text('allowed')
    cat.scan();drain(cat)
    hits=cat.search('strategy')['results'];assert len(hits)==1 and 'Downloads-extra' in hits[0]['path']
    with pytest.raises(ValueError):cat.read_bytes(cat.root/'Documents/alias.md')

def test_future_changes_deletion_duplicates_and_stale_versions(cat):
    path=cat.root/'Documents/strategy.md';path.write_text('Strategy roadmap.')
    cat.change(path);drain(cat);assert cat.search('strategy')['results']
    copy=cat.root/'Documents/copy.md';copy.write_bytes(path.read_bytes());cat.change(copy);drain(cat)
    hits=cat.search('strategy')['results'];assert len(hits)==1 and len(hits[0]['copies'])==1
    with cat.db() as db:assert db.execute('SELECT count(*) FROM documents').fetchone()[0]==1
    path.write_text('Changed financial budget');assert all(h['path']!=str(path) for h in cat.search('strategy')['results'])
    cat.change(path);drain(cat);assert cat.search('bütçe')['results'][0]['path']==str(path)
    copy.unlink();cat.change(copy);assert not cat.search('roadmap')['results']

def test_docx_full_extraction_and_metadata_only(cat):
    from docx import Document
    doc=Document();doc.add_paragraph('Strategic business plan');table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='Budget 2027'
    doc.save(cat.root/'Documents/unnamed.docx');(cat.root/'Documents/strategy.bin').write_bytes(b'\0\0binary')
    cat.scan();drain(cat)
    assert cat.search('bütçe docx')['results'][0]['name']=='unnamed.docx'
    assert cat.status()['counts']['metadata_only']==1

def test_restart_preserves_catalog(cat):
    p=cat.root/'Documents/strategy.md';p.write_text('Strategy');cat.scan();drain(cat)
    again=Catalog(cat.state,cat.config);assert again.search('strateji')['results']

def test_recent_edit_gets_priority_without_starving_backlog(cat):
    for i in range(12):(cat.root/f'Documents/old-{i}.md').write_text('old content '+str(i))
    cat.scan()
    new=cat.root/'Documents/new.md';new.write_text('fresh strategy content');cat.change(new)
    cat.process_one()
    with cat.db() as db:assert db.execute('SELECT status FROM files WHERE path=?',(str(new),)).fetchone()[0]=='indexed'
    drain(cat);assert cat.status()['counts']['indexed']==13

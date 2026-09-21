"""Resource-bounded local extraction, including native macOS Vision OCR."""
import io,os,re,sys,resource,subprocess,zipfile
from pathlib import Path
from xml.etree import ElementTree

resource.setrlimit(resource.RLIMIT_CPU,(220,220))
try:resource.setrlimit(resource.RLIMIT_DATA,(2*1024**3,2*1024**3))
except (ValueError,OSError):pass
max_bytes=int(sys.argv[2]) if len(sys.argv)>2 else 256*1024**2
ext=sys.argv[1];blob=sys.stdin.buffer.read(max_bytes+1)
try:
    if len(blob)>max_bytes:raise ValueError('document_limit')
    if ext in ('pdf','png','jpg','jpeg','tif','tiff','bmp','webp','heic'):
        ocr=Path(os.environ.get('LOCAL_FILES_OCR_EXECUTABLE',str(Path(__file__).with_name('ocr'))))
        if ocr.is_file():
            result=subprocess.run([str(ocr),ext],input=blob,capture_output=True,timeout=210)
            if result.returncode:raise ValueError('document_parse_failed')
            value=result.stdout.decode('utf-8')
        elif ext=='pdf':
            from pypdf import PdfReader
            pages=[p.extract_text() or '' for p in PdfReader(io.BytesIO(blob)).pages]
            if any(not p.strip() for p in pages):raise ValueError('ocr_unavailable')
            value='\n\n'.join(pages)
        else:raise ValueError('ocr_unavailable')
    elif ext=='docx':
        from docx import Document
        doc=Document(io.BytesIO(blob));value='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    elif ext=='xlsx':
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(blob),read_only=True,data_only=False)
        parts=[];size=0
        for sheet in book:
            parts.append('# '+sheet.title)
            for row in sheet.iter_rows():
                line=' | '.join(f'{c.coordinate}: {c.value}' for c in row if c.value is not None)
                size+=len(line)
                if size>128*1024**2:raise ValueError('document_limit')
                if line:parts.append(line)
        book.close();value='\n'.join(parts)
    elif ext=='pptx':
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            files=[n for n in archive.namelist() if re.fullmatch(r'ppt/(?:slides/slide|notesSlides/notesSlide)\d+\.xml',n)]
            if sum(archive.getinfo(n).file_size for n in files)>128*1024**2:raise ValueError('document_limit')
            value='\n\n'.join(n+'\n'+'\n'.join(el.text or '' for el in ElementTree.fromstring(archive.read(n)).iter() if el.tag.endswith('}t')) for n in sorted(files,key=lambda x:(x.startswith('ppt/notes'),int(re.search(r'(\d+)\.xml',x)[1]))))
    else:raise ValueError('document_parse_failed')
    if len(value.encode())>128*1024**2:raise ValueError('document_limit')
    if not re.sub(r'\[Page \d+\]','',value).strip():raise ValueError('empty_document')
    sys.stdout.write(value)
except Exception as exc:
    reason=str(exc) if isinstance(exc,ValueError) and str(exc) in ('document_limit','ocr_unavailable','empty_document') else 'document_parse_failed'
    sys.stderr.write(reason);sys.exit(1)

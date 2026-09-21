import asyncio
import io
import json
from pathlib import Path
import sys
import uuid
import pytest
import pytest_asyncio
from docx import Document
sys.path.insert(0,str(Path(__file__).parents[1]))
from engine import Engine,ServiceError
from importers import extract
from store import Store
from workflow import STAGES,ancestors,citations,digest,prompt_for

class Provider:
    def __init__(self):self.calls=[];self.fail=set();self.gates={}
    async def synthesize(self,stage,prompt):
        self.calls.append((stage['id'],prompt))
        if stage['id'] in self.gates:await self.gates[stage['id']].wait()
        if stage['id'] in self.fail:raise ServiceError('Test provider failure',502)
        await asyncio.sleep(.005)
        return 'Synthesis '+stage['id']+' with evidence https://example.org/source',{'total_tokens':15}
class Sink:
    def __init__(self):self.notes={}
    async def notebook(self,run):return 'notebook:test'
    async def note(self,run,stage):
        key=(run['id'],stage['id']);self.notes.setdefault(key,'note:'+stage['id']);return self.notes[key]
@pytest_asyncio.fixture
async def engine(tmp_path):
    store=Store(tmp_path);await store.open();provider=Provider();sink=Sink()
    engine=Engine(store,provider,sink,lambda text:len(text)//4,90000)
    yield engine
    await engine.close();await store.close()
async def create(engine,auto=True):
    return await engine.create({'question':'Which decision is supported by the evidence?','scope':'Compare independent primary sources','language':'English','auto_synthesize':auto,'notebook_id':None},uuid.uuid4().hex)
async def add(engine,run_id,stage_id,text=None):
    packet=await engine.packet(run_id,stage_id)
    return await engine.import_report(run_id,stage_id,text or ('Unique report '+stage_id+' with full evidence https://example.org/'+stage_id),[], '',packet['sha256'],[])
async def settle(engine):
    for _ in range(100):
        if not engine.tasks and not engine.background:return
        await asyncio.sleep(.01)
    raise AssertionError('Tasks did not settle')
@pytest.mark.asyncio
async def test_exact_eight_stages_and_round_barriers(engine):
    run=await create(engine);rid=run['id'];assert len(run['stages'])==8
    with pytest.raises(ServiceError):await add(engine,rid,'review_chatgpt')
    for sid in ['research_gemini','research_chatgpt']:await add(engine,rid,sid)
    assert not engine.provider.calls
    await add(engine,rid,'research_claude')
    await add(engine,rid,'review_chatgpt');assert not engine.provider.calls
    await add(engine,rid,'review_claude');await settle(engine)
    run=await engine.get(rid);assert run['status']=='completed'
    assert [x[0] for x in engine.provider.calls].count('final_chatgpt')==1
    assert len(engine.provider.calls)==3
    assert len(engine.sink.notes)==8
@pytest.mark.asyncio
async def test_first_round_is_independent_and_all_evidence_survives(engine):
    run=await create(engine,False);rid=run['id']
    await add(engine,rid,'research_gemini','Unique Gemini material '+('long evidence ' * 3000))
    other=await engine.packet(rid,'research_chatgpt');assert 'Unique Gemini material' not in other['prompt']
    await add(engine,rid,'research_chatgpt');await add(engine,rid,'research_claude')
    p1=await engine.packet(rid,'review_chatgpt');p2=await engine.packet(rid,'review_claude')
    # Same complete evidence, independent stage-specific claim ID namespaces.
    assert p1['prompt'].replace('review_chatgpt:C001','STAGE:C001')==p2['prompt'].replace('review_claude:C001','STAGE:C001')
    assert ('long evidence '*3000) in p1['prompt']
    assert p1['report_count']==3
@pytest.mark.asyncio
async def test_final_waits_for_both_syntheses(engine):
    gate=asyncio.Event();engine.provider.gates['synthesis_claude']=gate
    run=await create(engine)
    for sid,_,rnd,_ in STAGES:
        if rnd<=2:await add(engine,run['id'],sid)
    await asyncio.sleep(.08)
    assert 'final_chatgpt' not in [c[0] for c in engine.provider.calls]
    gate.set();await settle(engine)
    final=next(p for sid,p in engine.provider.calls if sid=='final_chatgpt')
    for sid,_,_,_ in STAGES[:-1]:assert sid in final
@pytest.mark.asyncio
async def test_idempotent_create_import_and_notebook_sync(engine):
    data={'question':'Question here','scope':'','language':'English','auto_synthesize':False,'notebook_id':None}
    a=await engine.create(data,'same-key-123');b=await engine.create(data,'same-key-123');assert a['id']==b['id']
    await add(engine,a['id'],'research_gemini');await add(engine,a['id'],'research_gemini');await engine.sync(a['id']);await engine.sync(a['id'])
    assert len(engine.sink.notes)==1
    with pytest.raises(ServiceError):await add(engine,a['id'],'research_gemini','Different report that cannot replace the completed original')
    with pytest.raises(ServiceError):await engine.create({**data,'question':'different'},'same-key-123')
@pytest.mark.asyncio
async def test_failed_sibling_does_not_trigger_final_or_repeat_completed(engine):
    engine.provider.fail.add('synthesis_claude');run=await create(engine)
    for sid,_,rnd,_ in STAGES:
        if rnd<=2:await add(engine,run['id'],sid)
    await settle(engine);assert (await engine.get(run['id']))['status']=='needs_attention'
    assert len(engine.provider.calls)==2
    engine.provider.fail.clear();await engine.action(run['id'],'resume');await settle(engine)
    names=[s for s,p in engine.provider.calls]
    assert names.count('synthesis_chatgpt')==1 and names.count('synthesis_claude')==2 and names.count('final_chatgpt')==1
@pytest.mark.asyncio
async def test_context_limit_preserves_full_packet_and_makes_no_model_call(engine):
    engine.token_limit=20;run=await create(engine)
    for sid,_,rnd,_ in STAGES:
        if rnd<=2:await add(engine,run['id'],sid)
    await settle(engine);assert not engine.provider.calls
    assert any(s['status']=='context_limit' for s in (await engine.get(run['id']))['stages'])
    assert (await engine.packet(run['id'],'synthesis_chatgpt'))['report_count']==5

@pytest.mark.asyncio
async def test_retry_uses_saved_full_input_and_rejects_tampering(engine):
    engine.provider.fail.add('synthesis_claude');run=await create(engine)
    for sid,_,rnd,_ in STAGES:
        if rnd<=2:await add(engine,run['id'],sid)
    await settle(engine)
    original=next(p for sid,p in engine.provider.calls if sid=='synthesis_claude')
    run=await engine.get(run['id']);stage=engine.stage(run,'synthesis_claude')
    path=engine.input_path(run,stage);assert path.read_text()==original
    run['stages'][0]['report']['content']+=' Presentation corrected after submission.'
    await engine.store.save(run)
    assert (await engine.packet(run['id'],stage['id']))['prompt']==original
    path.write_text('tampered')
    engine.provider.fail.clear();await engine.action(run['id'],'resume');await settle(engine)
    assert len(engine.provider.calls)==2
    assert engine.stage(await engine.get(run['id']),stage['id'])['status']=='failed'
    path.write_text(original)
    await engine.action(run['id'],'resume');await settle(engine)
    retried=[p for sid,p in engine.provider.calls if sid==stage['id']]
    assert retried==[original,original]
@pytest.mark.asyncio
async def test_restart_marks_inflight_interrupted_without_repeating(engine):
    run=await create(engine,False);run['stages'][0]['status']='completed';run['stages'][0]['report']={'content':'saved'}
    run['stages'][5]['status']='running';await engine.store.save(run);await engine.recover()
    restored=await engine.get(run['id']);assert restored['stages'][0]['report']['content']=='saved'
    assert restored['stages'][5]['status']=='interrupted';assert not engine.provider.calls
@pytest.mark.asyncio
async def test_pause_prevents_automatic_synthesis(engine):
    run=await create(engine);await engine.action(run['id'],'pause')
    for sid,_,rnd,_ in STAGES:
        if rnd<=2:await add(engine,run['id'],sid)
    assert not engine.provider.calls
    await engine.action(run['id'],'resume');await settle(engine);assert len(engine.provider.calls)==3
@pytest.mark.asyncio
async def test_stale_packet_rejected_and_evidence_preserved(engine):
    run=await create(engine,False);sid='research_gemini';packet=await engine.packet(run['id'],sid)
    with pytest.raises(ServiceError):await engine.import_report(run['id'],sid,'Some sufficiently long report text',[],'','wrong',[])
    ev=[{'name':'primary.txt','content':'PRIMARY SOURCE EXACT CONTENT https://example.org/evidence','sha256':'test'}]
    await engine.import_report(run['id'],sid,'Sufficiently long report text with supporting evidence',ev,'',packet['sha256'],[('original.txt',b'original bytes')])
    result=await engine.get(run['id']);assert result['stages'][0]['report']['evidence']==ev
    assert (engine.store.root/'artifacts'/run['id']/sid/'0.txt').read_bytes()==b'original bytes'
def test_document_extraction_and_invalid_files():
    assert extract('report.md','Türkçe rapor'.encode())=='Türkçe rapor'
    d=Document();d.add_paragraph('Research evidence');d.add_table(rows=1,cols=2).cell(0,0).text='Table evidence';buf=io.BytesIO();d.save(buf)
    result=extract('report.docx',buf.getvalue());assert 'Research evidence' in result and 'Table evidence' in result
    with pytest.raises(ValueError):extract('evil.exe',b'whatever')
    with pytest.raises(ValueError):extract('empty.txt',b'')

def test_source_links_keep_balanced_parentheses_and_remove_markdown_delimiter():
    assert citations('[Source](https://example.org/Report_(2026)). https://example.org/Report_(2026)') == ['https://example.org/Report_(2026)']

@pytest.mark.asyncio
async def test_dates_preserve_unknowns_and_research_independence(engine):
    run=await engine.create({'question':'Which current evidence supports this decision?', 'scope':'Primary sources', 'language':'Türkçe', 'as_of':'2026-09-21', 'auto_synthesize':False},uuid.uuid4().hex)
    rid=run['id'];packet=await engine.packet(rid,'research_gemini')
    await engine.import_report(rid,'research_gemini','Report with older research date and primary-source evidence https://example.org',[],'',packet['sha256'],[],'2026-08-01')
    await add(engine,rid,'research_chatgpt');await add(engine,rid,'research_claude')
    prompt=(await engine.packet(rid,'review_chatgpt'))['prompt']
    assert '2026-09-21' in prompt and '2026-08-01' in prompt
    assert '"researched_at": null' in prompt
    assert 'oy çokluğu' in prompt and 'kanıt tablosu' in prompt

@pytest.mark.asyncio
async def test_pause_prevents_queued_provider_request(engine):
    run=await create(engine)
    await engine.provider_locks['ChatGPT'].acquire()
    for sid,_,rnd,_ in STAGES:
        if rnd<=2:await add(engine,run['id'],sid)
    await engine.action(run['id'],'pause')
    engine.provider_locks['ChatGPT'].release()
    await settle(engine)
    assert not any(sid=='synthesis_chatgpt' for sid,_ in engine.provider.calls)
    assert (await engine.get(run['id']))['stages'][5]['status']=='ready'

@pytest.mark.asyncio
async def test_only_unstarted_import_run_can_be_automated(engine,monkeypatch):
    run=await create(engine)
    kicked=[]
    async def kick(rid):kicked.append(rid)
    monkeypatch.setattr(engine,'kick',kick)
    with pytest.raises(ServiceError):await engine.action(run['id'],'automate')
    engine.browser=object()
    converted=await engine.action(run['id'],'automate')
    assert converted['id']==run['id'] and converted['execution_mode']=='browser'
    assert [s['mode'] for s in converted['stages']]==['browser']*5+['account']*3
    assert kicked==[run['id']]
    await engine.action(run['id'],'automate')
    assert kicked==[run['id']]
    second=await create(engine,False)
    await add(engine,second['id'],'research_gemini')
    with pytest.raises(ServiceError):await engine.action(second['id'],'automate')
    engine.browser=None
